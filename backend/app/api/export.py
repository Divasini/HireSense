import csv
import io
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.job import Job
from app.models.application import Application, CandidateStatus
from app.models.candidate import Candidate
from app.models.screening import ScreeningResult

router = APIRouter(prefix="/api/export", tags=["export"])

@router.get("/candidates/csv/{job_id}")
def export_csv(job_id: str, db: Session = Depends(get_db)):
    job = db.query(Job).filter(Job.id == job_id).first()
    apps = db.query(Application).filter(Application.job_id == job_id).all()
    
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Rank", "Candidate Name", "Email", "Overall Score (%)", "Skills Score", "Experience Score", "Recommendation", "Status"])
    
    items = []
    for a in apps:
        cand = db.query(Candidate).filter(Candidate.id == a.candidate_id).first()
        scr = db.query(ScreeningResult).filter(ScreeningResult.application_id == a.id).first()
        stat = db.query(CandidateStatus).filter(CandidateStatus.application_id == a.id).first()
        items.append({
            "name": cand.full_name if cand else "Unknown",
            "email": cand.email if cand else "",
            "score": scr.overall_score if scr else 0.0,
            "skills": scr.skills_score if scr else 0.0,
            "exp": scr.experience_score if scr else 0.0,
            "rec": stat.ai_recommendation if stat else "consider",
            "status": stat.recruiter_decision if (stat and stat.recruiter_decision) else a.status
        })
        
    items = sorted(items, key=lambda x: x["score"], reverse=True)
    for idx, item in enumerate(items):
        writer.writerow([idx + 1, item["name"], item["email"], item["score"], item["skills"], item["exp"], item["rec"], item["status"]])
        
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=candidates_{job_id}.csv"}
    )

@router.get("/candidates/excel/{job_id}")
def export_excel(job_id: str, db: Session = Depends(get_db)):
    import openpyxl
    job = db.query(Job).filter(Job.id == job_id).first()
    apps = db.query(Application).filter(Application.job_id == job_id).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screening Results"
    ws.append(["Rank", "Candidate Name", "Email", "Overall Score (%)", "Skills Score", "Experience Score", "Recommendation", "Status"])
    
    items = []
    for a in apps:
        cand = db.query(Candidate).filter(Candidate.id == a.candidate_id).first()
        scr = db.query(ScreeningResult).filter(ScreeningResult.application_id == a.id).first()
        stat = db.query(CandidateStatus).filter(CandidateStatus.application_id == a.id).first()
        items.append({
            "name": cand.full_name if cand else "Unknown",
            "email": cand.email if cand else "",
            "score": scr.overall_score if scr else 0.0,
            "skills": scr.skills_score if scr else 0.0,
            "exp": scr.experience_score if scr else 0.0,
            "rec": stat.ai_recommendation if stat else "consider",
            "status": stat.recruiter_decision if (stat and stat.recruiter_decision) else a.status
        })
        
    items = sorted(items, key=lambda x: x["score"], reverse=True)
    for idx, item in enumerate(items):
        ws.append([idx + 1, item["name"], item["email"], item["score"], item["skills"], item["exp"], item["rec"], item["status"]])
        
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=candidates_{job_id}.xlsx"}
    )

@router.get("/screening/pdf/{application_id}")
def export_pdf(application_id: str, db: Session = Depends(get_db)):
    from fpdf import FPDF
    app = db.query(Application).filter(Application.id == application_id).first()
    if not app:
        raise HTTPException(status_code=404, detail="Application not found")
        
    cand = db.query(Candidate).filter(Candidate.id == app.candidate_id).first()
    job = db.query(Job).filter(Job.id == app.job_id).first()
    scr = db.query(ScreeningResult).filter(ScreeningResult.application_id == application_id).first()
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "AI Candidate Screening Report", ln=True, align="C")
    pdf.ln(5)
    
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, f"Candidate: {cand.full_name if cand else 'Unknown'}", ln=True)
    pdf.cell(0, 8, f"Email: {cand.email if cand else ''}", ln=True)
    pdf.cell(0, 8, f"Applied Role: {job.title if job else 'Position'}", ln=True)
    pdf.ln(5)
    
    if scr:
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, f"Overall Match Score: {scr.overall_score}%", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 7, f"- Skills Score: {scr.skills_score}/100", ln=True)
        pdf.cell(0, 7, f"- Experience Score: {scr.experience_score}/100", ln=True)
        pdf.cell(0, 7, f"- Education Score: {scr.education_score}/100", ln=True)
        pdf.cell(0, 7, f"- Project Relevance: {scr.project_score}/100", ln=True)
        pdf.cell(0, 7, f"- Certifications: {scr.certification_score}/100", ln=True)
        pdf.ln(5)
        
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "AI Explainable Summary:", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, scr.ai_summary or scr.explanation or "Matches primary job criteria.")
    
    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=screening_{application_id}.pdf"}
    )
